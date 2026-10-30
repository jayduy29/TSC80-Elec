"""
Script đồng bộ dữ liệu hàng tuần từ file TCS80-C459-ELECTRICAL PROGRESS.xlsx
vào C459_ELECTRICAL_PROGRESS_OVERVIEW.xlsx, prog_data.json, và index.html (D.cbl).

Sử dụng:
  python sync_progress.py                  # Tự động nhận diện tuần hoặc dùng tuần hiện tại
  python sync_progress.py --week W44       # Chỉ định tuần W44
  python sync_progress.py --no-push        # Không tự động git push
"""

import os
import sys
import json
import re
import subprocess
from datetime import date
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

MASTER_FILE = 'TCS80-C459-ELECTRICAL PROGRESS.xlsx'
OVERVIEW_FILE = 'C459_ELECTRICAL_PROGRESS_OVERVIEW.xlsx'
JSON_OUT = 'prog_data.json'
HTML_FILE = 'index.html'

CONTR_PLANS = {
    'HDC': {'cut': 30259, 'pull': 30259, 'conn': 2694},
    'NAUTICS': {'cut': 26381, 'pull': 26381, 'conn': 2006},
    'ST': {'cut': 8593, 'pull': 8593, 'conn': 554},
    'TOTAL': {'cut': 65233, 'pull': 65233, 'conn': 5254}
}

ROW_MAP = {
    # contractor: {item: row_idx}
    'HDC': {'cut': 7, 'pull': 9, 'conn': 11},
    'NAUTICS': {'cut': 13, 'pull': 15, 'conn': 17},
    'ST': {'cut': 19, 'pull': 21, 'conn': 23},
    'TOTAL': {'cut': 25, 'pull': 27, 'conn': 29}
}

ROW_PLANS = {
    7: 30259, 9: 30259, 11: 2694,
    13: 26381, 15: 26381, 17: 2006,
    19: 8593, 21: 8593, 23: 554,
    25: 65233, 27: 65233, 29: 5254
}

def to_num(val):
    if val is None or val == '' or str(val).strip() in ('-', '#VALUE!', '#N/A'):
        return 0
    try:
        return float(str(val).replace(',', ''))
    except:
        return 0

def to_int(val, default=0):
    if val is None: return default
    try:
        return int(float(str(val).strip()))
    except:
        return default

def read_tong_sheet(excel_path):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if 'TỔNG' not in wb.sheetnames:
        raise ValueError(f"Sheet 'TỔNG' không tồn tại trong {excel_path}")
    ws = wb['TỔNG']

    data = {
        'HDC': {'cut': 0, 'pull': 0, 'conn': 0},
        'NAUTICS': {'cut': 0, 'pull': 0, 'conn': 0},
        'ST': {'cut': 0, 'pull': 0, 'conn': 0},
        'TOTAL': {'cut': 0, 'pull': 0, 'conn': 0}
    }

    for r in range(1, ws.max_row + 1):
        c1 = str(ws.cell(r, 1).value or '').strip().lower()
        if 'total' in c1 or 'tổng' in c1:
            continue
        if 'cut' in c1 or 'cắt' in c1:
            data['HDC']['cut'] = to_num(ws_tong_val(ws, r, 2))
            data['NAUTICS']['cut'] = to_num(ws_tong_val(ws, r, 4))
            data['ST']['cut'] = to_num(ws_tong_val(ws, r, 6))
            data['TOTAL']['cut'] = to_num(ws_tong_val(ws, r, 8))
        elif 'pull' in c1 or 'kéo' in c1:
            data['HDC']['pull'] = to_num(ws_tong_val(ws, r, 2))
            data['NAUTICS']['pull'] = to_num(ws_tong_val(ws, r, 4))
            data['ST']['pull'] = to_num(ws_tong_val(ws, r, 6))
            data['TOTAL']['pull'] = to_num(ws_tong_val(ws, r, 8))
        elif 'connect' in c1 or 'đấu' in c1:
            data['HDC']['conn'] = to_num(ws_tong_val(ws, r, 2))
            data['NAUTICS']['conn'] = to_num(ws_tong_val(ws, r, 4))
            data['ST']['conn'] = to_num(ws_tong_val(ws, r, 6))
            data['TOTAL']['conn'] = to_num(ws_tong_val(ws, r, 8))

    return data

def ws_tong_val(ws, r, c):
    v = ws.cell(r, c).value
    return v

def update_overview_excel(data, target_week):
    if not os.path.exists(OVERVIEW_FILE):
        raise FileNotFoundError(f"{OVERVIEW_FILE} không tìm thấy!")

    wb_ov = openpyxl.load_workbook(OVERVIEW_FILE)
    col_found = -1

    for sname in wb_ov.sheetnames:
        ws = wb_ov[sname]
        col_idx = -1
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(5, c).value or '').strip()
            if val.upper() == target_week.upper():
                col_idx = c
                break

        if col_idx == -1:
            print(f"⚠️ Cảnh báo: Không tìm thấy cột {target_week} trong sheet {sname}")
            continue

        col_found = col_idx

        # Update values
        for contr, items in data.items():
            for item_k, val in items.items():
                r = ROW_MAP[contr][item_k]
                ws.cell(r, col_idx).value = int(val) if val else 0

        # Totals
        total_cut = data['HDC']['cut'] + data['NAUTICS']['cut'] + data['ST']['cut']
        total_pull = data['HDC']['pull'] + data['NAUTICS']['pull'] + data['ST']['pull']
        total_conn = data['HDC']['conn'] + data['NAUTICS']['conn'] + data['ST']['conn']

        ws.cell(25, col_idx).value = int(total_cut)
        ws.cell(27, col_idx).value = int(total_pull)
        ws.cell(29, col_idx).value = int(total_conn)

    wb_ov.save(OVERVIEW_FILE)
    print(f"✅ Đã cập nhật {OVERVIEW_FILE} (cột {target_week})")
    return col_found

def generate_prog_json(target_week, col_idx, data):
    wb_data = openpyxl.load_workbook(OVERVIEW_FILE, data_only=True)
    sheets_json = []

    total_cut = data['HDC']['cut'] + data['NAUTICS']['cut'] + data['ST']['cut']
    total_pull = data['HDC']['pull'] + data['NAUTICS']['pull'] + data['ST']['pull']
    total_conn = data['HDC']['conn'] + data['NAUTICS']['conn'] + data['ST']['conn']

    for sname in wb_data.sheetnames:
        ws = wb_data[sname]
        rows = []
        for r in range(1, ws.max_row + 1):
            row_vals = []
            for c in range(1, ws.max_column + 1):
                val = ws.cell(r, c).value
                if c == col_idx:
                    if r == 7: val = int(data['HDC']['cut'])
                    elif r == 9: val = int(data['HDC']['pull'])
                    elif r == 11: val = int(data['HDC']['conn'])
                    elif r == 13: val = int(data['NAUTICS']['cut'])
                    elif r == 15: val = int(data['NAUTICS']['pull'])
                    elif r == 17: val = int(data['NAUTICS']['conn'])
                    elif r == 19: val = int(data['ST']['cut'])
                    elif r == 21: val = int(data['ST']['pull'])
                    elif r == 23: val = int(data['ST']['conn'])
                    elif r == 25: val = int(total_cut)
                    elif r == 27: val = int(total_pull)
                    elif r == 29: val = int(total_conn)
                    elif r in (8, 10, 12, 14, 16, 18, 20, 22, 24, 26, 28, 30):
                        parent_r = r - 1
                        if parent_r == 7: pval = data['HDC']['cut']
                        elif parent_r == 9: pval = data['HDC']['pull']
                        elif parent_r == 11: pval = data['HDC']['conn']
                        elif parent_r == 13: pval = data['NAUTICS']['cut']
                        elif parent_r == 15: pval = data['NAUTICS']['pull']
                        elif parent_r == 17: pval = data['NAUTICS']['conn']
                        elif parent_r == 19: pval = data['ST']['cut']
                        elif parent_r == 21: pval = data['ST']['pull']
                        elif parent_r == 23: pval = data['ST']['conn']
                        elif parent_r == 25: pval = total_cut
                        elif parent_r == 27: pval = total_pull
                        elif parent_r == 29: pval = total_conn
                        else: pval = 0
                        try:
                            val = float(pval) / ROW_PLANS[parent_r]
                        except:
                            val = ''

                if val is None:
                    row_vals.append('')
                elif isinstance(val, float):
                    if val.is_integer():
                        row_vals.append(str(int(val)))
                    else:
                        row_vals.append(f"{val:.6g}")
                else:
                    row_vals.append(str(val))
            if any(x != '' for x in row_vals):
                rows.append(row_vals)
        sheets_json.append({"name": sname, "rows": rows})

    with open(JSON_OUT, 'w', encoding='utf-8') as f:
        json.dump(sheets_json, f, ensure_ascii=False, indent=2)

    print(f"✅ Đã tạo {JSON_OUT} ({len(sheets_json)} sheets)")

def sync_cables_to_html(excel_path):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if 'Consolidated' not in wb.sheetnames:
        print("⚠️ Sheet 'Consolidated' không có, bỏ qua đồng bộ cáp vào index.html")
        return 0

    ws = wb['Consolidated']
    header = [str(ws.cell(1, c).value or '').strip() for c in range(1, 25)]

    col_contr = header.index('Contractor') + 1
    col_cid = header.index('Cable ID') + 1
    col_cut = header.index('Cut') + 1
    col_pulled = header.index('Pulled') + 1
    col_conn = header.index('Connection') + 1
    col_megger = header.index('CWC- Megger') + 1
    col_fa = header.index('From Device Area') + 1
    col_ta = header.index('To Device Area') + 1

    contr_map = {'HDC': 0, 'NAUTICS': 1, 'ST': 2}

    new_cbl = {}
    for r in range(2, ws.max_row + 1):
        cid = ws.cell(r, col_cid).value
        if not cid: continue
        cid = str(cid).strip()

        c_name = str(ws.cell(r, col_contr).value or '').strip().upper()
        c_idx = contr_map.get(c_name, 3)

        cut_v = to_int(ws.cell(r, col_cut).value, 0)
        pulled_v = to_int(ws.cell(r, col_pulled).value, 0)
        conn_v = to_int(ws.cell(r, col_conn).value, 0)
        megger_v = to_int(ws.cell(r, col_megger).value, 0)

        fa = str(ws.cell(r, col_fa).value or '').strip()
        ta = str(ws.cell(r, col_ta).value or '').strip()

        new_cbl[cid] = [c_idx, pulled_v, cut_v, conn_v, fa, ta, megger_v]

    with open(HTML_FILE, 'r', encoding='utf-8') as f:
        html = f.read()

    pos1 = html.find('"cbl":{')
    if pos1 == -1:
        print("⚠️ Không tìm thấy \"cbl\":{ trong index.html")
        return 0
    pos2 = html.find('}};', pos1)
    if pos2 == -1:
        print("⚠️ Không tìm thấy kết thúc cbl trong index.html")
        return 0

    cur_cbl = json.loads(html[pos1+6:pos2+1])
    diff_count = 0
    for cid, nval in new_cbl.items():
        if cur_cbl.get(cid) != nval:
            diff_count += 1
            cur_cbl[cid] = nval

    if diff_count > 0:
        new_cbl_str = json.dumps(cur_cbl, separators=(',', ':'))
        new_html = html[:pos1+6] + new_cbl_str[1:-1] + html[pos2:]
        with open(HTML_FILE, 'w', encoding='utf-8') as f:
            f.write(new_html)
        print(f"✅ Đã cập nhật {diff_count} dây thay đổi trạng thái vào {HTML_FILE}")
    else:
        print(f"ℹ️ Trạng thái cáp trên bản vẽ {HTML_FILE} đã hoàn toàn khớp (0 thay đổi)")

    return diff_count

def git_commit_and_push(target_week):
    try:
        subprocess.run(['git', 'add', OVERVIEW_FILE, JSON_OUT, HTML_FILE, MASTER_FILE], check=True)
        msg = f"feat(prog): cap nhat tien do tuan {target_week} tu {MASTER_FILE}"
        subprocess.run(['git', 'commit', '-m', msg], check=False)
        res = subprocess.run(['git', 'push', 'origin', 'main'], capture_output=True, text=True)
        if res.returncode == 0:
            print(f"🚀 Đã tự động đẩy lên GitHub (nhánh main)! Web và điện thoại sẽ nhận bản mới sau 30s.")
        else:
            print("⚠️ Git push gặp lỗi (có thể kiểm tra mạng):", res.stderr)
    except Exception as e:
        print("⚠️ Lỗi git commit/push:", e)

def main():
    import argparse
    parser = argparse.ArgumentParser(description="Đồng bộ tiến độ hàng tuần từ Excel vào web")
    parser.add_argument('--file', default=MASTER_FILE, help="Đường dẫn file Excel tiến độ tuần")
    parser.add_argument('--week', default=None, help="Mã tuần (ví dụ W44, W45). Mặc định tự tính.")
    parser.add_argument('--no-push', action='store_true', help="Không tự động git push")
    args = parser.parse_args()

    file_path = args.file
    if not os.path.exists(file_path):
        print(f"❌ Không tìm thấy file {file_path}")
        sys.exit(1)

    print(f"📂 Đang đọc dữ liệu từ: {file_path}")
    data = read_tong_sheet(file_path)

    print("\n--- SỐ LIỆU TỔNG HỢP THỰC HIỆN ---")
    for k in ['HDC', 'NAUTICS', 'ST', 'TOTAL']:
        v = data[k]
        print(f"  {k:12} | Cắt: {v['cut']:8,g} m | Kéo: {v['pull']:8,g} m | Đấu: {v['conn']:5,g}")

    target_week = args.week
    if not target_week:
        # Check if filename has W\d\d
        m = re.search(r'W(\d{2})', file_path, re.I)
        if m:
            target_week = f"W{m.group(1)}"
        else:
            today = date.today()
            target_week = f"W{today.isocalendar()[1]:02d}"

    print(f"\n🎯 Đồng bộ cho tuần: {target_week}")

    col_idx = update_overview_excel(data, target_week)
    if col_idx != -1:
        generate_prog_json(target_week, col_idx, data)

    sync_cables_to_html(file_path)

    if not args.no_push:
        print("\n📤 Đang đẩy thay đổi lên GitHub...")
        git_commit_and_push(target_week)

    print(f"\n🎉 HOÀN THÀNH ĐỒNG BỘ TUẦN {target_week}!")

if __name__ == '__main__':
    main()
